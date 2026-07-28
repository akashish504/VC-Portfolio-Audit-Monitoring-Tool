"""XLSX export for the Timeline dashboard tables.

Each Timeline table can be downloaded as the grid the UI shows (strategy counts,
audit-status breakdown, and the FYE-month cross-tabs). The aggregated figures come
straight from the provider's ``timeline_summary`` so the workbook always matches the
on-screen numbers.
"""
from __future__ import annotations

import io
from datetime import date
from typing import Any

from sqlalchemy.ext.asyncio import AsyncSession

from src.services.dashboard import get_dashboard_provider

# Valid ``table`` selectors (the UI passes one per download button); "all" = every sheet.
_TABLES = ("strategy_all", "strategy_included", "analysis", "status", "matrices", "all")


def _style(ws, header_row: int, ncols: int) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill

    for ci in range(1, ncols + 1):
        cell = ws.cell(row=header_row, column=ci)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1D4ED8")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _autosize(ws, ncols: int) -> None:
    from openpyxl.utils import get_column_letter

    for ci in range(1, ncols + 1):
        width = 12
        for cell in ws[get_column_letter(ci)]:
            if cell.value is not None:
                width = max(width, len(str(cell.value)) + 4)
        ws.column_dimensions[get_column_letter(ci)].width = min(40, width)


def _write_strategy(ws, table) -> None:
    ws.append(["Strategy", "India", "SEA", "Total"])
    _style(ws, 1, 4)
    for r in table.rows:
        ws.append([r.strategy, r.india, r.sea, r.total])
    t = table.total
    ws.append([t.strategy or "Total", t.india, t.sea, t.total])
    _autosize(ws, 4)


def _write_analysis(ws, summary) -> None:
    ws.append(["Cohort", "India", "SEA", "Total"])
    _style(ws, 1, 4)
    a, i = summary.analysis_all, summary.analysis_included
    ws.append(["All companies", a.india, a.sea, a.total])
    ws.append(["Scoped in", i.india, i.sea, i.total])
    _autosize(ws, 4)


def _write_status(ws, table) -> None:
    strategies = list(table.strategies)
    headers = ["Status", "Overall", *strategies, "India", "SEA"]
    ws.append(headers)
    _style(ws, 1, len(headers))

    def row_vals(row) -> list[Any]:
        return [
            row.status,
            row.overall,
            *[(row.by_strategy or {}).get(s, 0) for s in strategies],
            row.india,
            row.sea,
        ]

    for row in table.rows:
        ws.append(row_vals(row))
    tot = table.total
    ws.append(["Total", tot.overall, *[(tot.by_strategy or {}).get(s, 0) for s in strategies], tot.india, tot.sea])
    _autosize(ws, len(headers))


def _write_matrices(ws, matrices) -> None:
    """All FYE-month cross-tabs stacked one after another (scope/geo labelled)."""
    row = 1
    for m in matrices:
        ws.cell(row=row, column=1, value=f"Audit timeline : {m.scope} — {m.geo}")
        from openpyxl.styles import Font

        ws.cell(row=row, column=1).font = Font(bold=True)
        row += 1
        headers = ["Status", *m.months, "Grand Total"]
        for ci, h in enumerate(headers, 1):
            ws.cell(row=row, column=ci, value=h)
        _style(ws, row, len(headers))
        row += 1
        for r in m.rows:
            ws.cell(row=row, column=1, value=r.status)
            for ci, mo in enumerate(m.months, 2):
                ws.cell(row=row, column=ci, value=(r.by_month or {}).get(mo, 0))
            ws.cell(row=row, column=len(headers), value=r.total)
            row += 1
        ws.cell(row=row, column=1, value="Grand Total")
        for ci, mo in enumerate(m.months, 2):
            ws.cell(row=row, column=ci, value=(m.column_totals or {}).get(mo, 0))
        ws.cell(row=row, column=len(headers), value=m.grand_total)
        row += 2  # blank spacer between matrices
    _autosize(ws, 16)


async def build_timeline_xlsx(
    db: AsyncSession, *, review_cycle_id: str | None = None, table: str = "all"
) -> tuple[bytes, str]:
    """Build the requested Timeline table(s) as an XLSX workbook."""
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - openpyxl is a hard dep in prod
        raise RuntimeError("openpyxl is not installed") from exc

    key = table if table in _TABLES else "all"
    summary = await get_dashboard_provider().timeline_summary(db, review_cycle_id=review_cycle_id)

    wb = openpyxl.Workbook()
    wb.remove(wb.active)  # start clean; add only the requested sheets

    if key in ("strategy_all", "all"):
        _write_strategy(wb.create_sheet("Strategy — All"), summary.strategy_all)
    if key in ("strategy_included", "all"):
        _write_strategy(wb.create_sheet("Strategy — Scoped in"), summary.strategy_included)
    if key in ("analysis", "all"):
        _write_analysis(wb.create_sheet("For analysis"), summary)
    if key in ("status", "all"):
        _write_status(wb.create_sheet("Audit status"), summary.status_table)
    if key in ("matrices", "all"):
        _write_matrices(wb.create_sheet("Audit timeline"), summary.matrices)

    if not wb.sheetnames:  # guard: never save an empty workbook
        wb.create_sheet("Timeline")

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    cycle = f"_{review_cycle_id}" if review_cycle_id else ""
    name = key if key != "all" else "timeline"
    return buf.read(), f"timeline_{name}{cycle}_{date.today()}.xlsx"
