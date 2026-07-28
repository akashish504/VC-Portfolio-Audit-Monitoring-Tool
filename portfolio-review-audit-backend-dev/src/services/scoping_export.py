"""XLSX exports for the scoping dashboard list views.

These export the COMPLETE result set (not the paginated page the UI shows) for a
given filter set, so a "Download" on any breakdown list — IL-wise, Category-wise,
or a drill-down — yields every matching row. The provider is queried with a very
large limit to bypass pagination.
"""
from __future__ import annotations

import io
from datetime import date
from typing import Any, Optional, Sequence

from sqlalchemy.ext.asyncio import AsyncSession

from src.services.dashboard import get_dashboard_provider

# Large enough to return the full set in one shot (company/deal counts are in the
# low thousands at most); avoids the route's paginated le=500 limit.
_EXPORT_LIMIT = 1_000_000


# (attribute, header) — order defines the sheet columns.
_COMPANY_COLUMNS: list[tuple[str, str]] = [
    ("company_id", "CID"),
    ("company_name", "Company"),
    ("fund", "Fund"),
    ("geography", "Geography"),
    ("fy_end", "FYE"),
    ("strategy", "Strategy"),
    ("investment_lead_1", "IL 1"),
    ("investment_lead_2", "IL 2"),
    ("consolidated_cost", "Cons. Cost"),
    ("consolidated_fmv", "Cons. FMV"),
    ("latest_category", "Category"),
    ("audit_status", "Audit Status"),
    ("review_stage", "Review Stage"),
    ("last_year_audit_status", "Last Yr Audit"),
]

_PCM_COLUMNS: list[tuple[str, str]] = [
    ("deal_name", "Deal Name"),
    ("fund", "Fund"),
    ("deal_id", "Deal ID"),
    ("strategy", "Strategy"),
    ("category", "Category"),
    ("fy_end", "FYE"),
    ("il_main", "IL"),
    ("deal_level_stage_1", "Stage 1"),
    ("deal_level_stage_2", "Stage 2"),
    ("auditor", "Auditor"),
    ("category_of_auditor", "Auditor Category"),
    ("consolidated_cost", "Cons. Cost"),
    ("consolidated_fmv", "Cons. FMV"),
    ("scoping_for_audit", "Scoped In"),
    ("reason_for_exclusion", "Reason for Exclusion"),
]


def _cell(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Yes" if value else "No"
    return value


def _rows_to_xlsx(title: str, columns: Sequence[tuple[str, str]], rows: Sequence[Any]) -> bytes:
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - openpyxl is a hard dep in prod
        raise RuntimeError("openpyxl is not installed") from exc

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title

    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1D4ED8")
    hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers = [label for _attr, label in columns]
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = hdr_align
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A2"

    for r in rows:
        ws.append([_cell(getattr(r, attr, None)) for attr, _label in columns])

    for ci, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(ci)].width = max(14, min(48, len(h) + 6))

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


async def build_scoping_companies_xlsx(db: AsyncSession, **filters: Any) -> tuple[bytes, str]:
    """Complete (non-paginated) company list for the given scoping filters."""
    items, _total = await get_dashboard_provider().scoping_companies(
        db, **filters, limit=_EXPORT_LIMIT, offset=0
    )
    xlsx = _rows_to_xlsx("Companies", _COMPANY_COLUMNS, items)
    cycle = filters.get("review_cycle_id")
    suffix = f"_{cycle}" if cycle else ""
    return xlsx, f"scoping_companies{suffix}_{date.today()}.xlsx"


async def build_pcm_deals_xlsx(
    db: AsyncSession,
    *,
    deal_level_stage_1: Optional[str] = None,
    deal_level_stage_2: Optional[str] = None,
    scoping_for_audit: Optional[bool] = None,
    fy_end: Optional[str] = None,
    category_of_auditor: Optional[str] = None,
    strategy: Optional[str] = None,
    geo_l1: Optional[str] = None,
    india_sea_only: bool = False,
    deal_ids: Optional[list[str]] = None,
    unique: bool = False,
    fund: Optional[str] = None,
    sector: Optional[str] = None,
    category: Optional[str] = None,
    investment_lead: Optional[str] = None,
    review_cycle_id: Optional[str] = None,
) -> tuple[bytes, str]:
    """Complete (non-paginated) master-scoping deal list for the given filters."""
    items, _total = await get_dashboard_provider().pcm_deals_by_stage(
        db,
        deal_level_stage_1=deal_level_stage_1,
        deal_level_stage_2=deal_level_stage_2,
        scoping_for_audit=scoping_for_audit,
        fy_end=fy_end,
        category_of_auditor=category_of_auditor,
        strategy=strategy,
        geo_l1=geo_l1,
        india_sea_only=india_sea_only,
        deal_ids=deal_ids,
        unique=unique,
        fund=fund,
        sector=sector,
        category=category,
        investment_lead=investment_lead,
        review_cycle_id=review_cycle_id,
        limit=_EXPORT_LIMIT,
        offset=0,
    )
    xlsx = _rows_to_xlsx("Deals", _PCM_COLUMNS, items)
    suffix = f"_{review_cycle_id}" if review_cycle_id else ""
    return xlsx, f"scoping_deals{suffix}_{date.today()}.xlsx"
