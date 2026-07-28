"""XLSX exports for the aggregated dashboard grids (Scoping + Variance).

These export the table *as shown* (the aggregated numbers), unlike the company/deal
list exports. Each builder re-queries the provider so the workbook always matches the
on-screen figures, then serialises the read-model into a simple header + rows grid.
"""
from __future__ import annotations

import io
from datetime import date
from typing import Any, Sequence

from sqlalchemy.ext.asyncio import AsyncSession

from src.services.dashboard import get_dashboard_provider


def _grid_to_xlsx(title: str, header: Sequence[Any], rows: Sequence[Sequence[Any]]) -> bytes:
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - openpyxl is a hard dep in prod
        raise RuntimeError("openpyxl is not installed") from exc

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append(list(header))
    for ci in range(1, len(header) + 1):
        c = ws.cell(row=1, column=ci)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1D4ED8")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    for r in rows:
        ws.append(["" if v is None else v for v in r])
    for ci in range(1, len(header) + 1):
        width = 12
        for cell in ws[get_column_letter(ci)]:
            if cell.value is not None:
                width = max(width, len(str(cell.value)) + 4)
        ws.column_dimensions[get_column_letter(ci)].width = min(44, width)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ── read-model → (header, rows) serialisers ──────────────────────────────────
def _matrix_grid(matrix) -> tuple[list, list]:
    header = ["Status", *matrix.months, "Total"]
    rows = [
        [r.status, *[(r.by_month or {}).get(m, 0) for m in matrix.months], r.total]
        for r in matrix.rows
    ]
    rows.append(["Total", *[(matrix.column_totals or {}).get(m, 0) for m in matrix.months], matrix.grand_total])
    return header, rows


def _yoy_grid(summary) -> tuple[list, list]:
    header = ["Category", summary.current_label or "Current", summary.prior_label or "Prior"]
    rows = [[r.category, r.current, r.prior] for r in summary.rows]
    rows.append(["Total", summary.current_total, summary.prior_total])
    return header, rows


def _discrepancy_summary_grid(summary) -> tuple[list, list]:
    header = ["Metric", "> threshold", "< threshold", "Not comparable", "Total"]
    rows = [[m.metric, m.gt_10, m.lt_10, m.not_comparable, m.total] for m in summary.metrics]
    return header, rows


def _reason_matrix_grid(matrix) -> tuple[list, list]:
    header = ["Metric", "Reason", "Count"]
    rows: list[list] = []
    for m in matrix.metrics:
        if not m.rows:
            rows.append([m.metric, "No discrepancies", 0])
            continue
        for r in m.rows:
            rows.append([m.metric, r.subcategory, r.count])
    return header, rows


def _suffix(review_cycle_id: str | None) -> str:
    return f"_{review_cycle_id}" if review_cycle_id else ""


# ── async builders (one per downloadable grid) ───────────────────────────────
async def build_status_matrix_xlsx(
    db: AsyncSession, *, view: str, review_cycle_id: str | None = None
) -> tuple[bytes, str]:
    prov = get_dashboard_provider()
    matrix = await (
        prov.scoping_actual_matrix(db, review_cycle_id=review_cycle_id)
        if view == "actual"
        else prov.scoping_tentative_matrix(db, review_cycle_id=review_cycle_id)
    )
    header, rows = _matrix_grid(matrix)
    return _grid_to_xlsx("Scoping status", header, rows), f"scoping_status_{view}{_suffix(review_cycle_id)}_{date.today()}.xlsx"


async def build_deal_stage_matrix_xlsx(
    db: AsyncSession, *, stage: str, basis: str, review_cycle_id: str | None = None, **dims
) -> tuple[bytes, str]:
    matrix = await get_dashboard_provider().deal_stage_matrix(
        db, review_cycle_id=review_cycle_id, stage=stage, basis=basis, **dims
    )
    header, rows = _matrix_grid(matrix)
    return _grid_to_xlsx(f"Deal stage {stage}", header, rows), f"deal_stage_{stage}{_suffix(review_cycle_id)}_{date.today()}.xlsx"


async def build_overall_status_xlsx(
    db: AsyncSession, *, review_cycle_id: str | None = None, **dims
) -> tuple[bytes, str]:
    summary = await get_dashboard_provider().scoping_overall_status(db, review_cycle_id=review_cycle_id, **dims)
    header, rows = _yoy_grid(summary)
    return _grid_to_xlsx("Overall status", header, rows), f"overall_status{_suffix(review_cycle_id)}_{date.today()}.xlsx"


async def build_auditor_summary_xlsx(
    db: AsyncSession, *, review_cycle_id: str | None = None, **dims
) -> tuple[bytes, str]:
    summary = await get_dashboard_provider().scoping_auditor_summary(db, review_cycle_id=review_cycle_id, **dims)
    header, rows = _yoy_grid(summary)
    return _grid_to_xlsx("Auditor summary", header, rows), f"auditor_summary{_suffix(review_cycle_id)}_{date.today()}.xlsx"


async def build_discrepancy_summary_xlsx(
    db: AsyncSession, *, review_cycle_id: str | None = None
) -> tuple[bytes, str]:
    summary = await get_dashboard_provider().variance_discrepancy_summary(db, review_cycle_id=review_cycle_id)
    header, rows = _discrepancy_summary_grid(summary)
    return _grid_to_xlsx("Discrepancy summary", header, rows), f"discrepancy_summary{_suffix(review_cycle_id)}_{date.today()}.xlsx"


async def build_discrepancy_reasons_xlsx(
    db: AsyncSession, *, review_cycle_id: str | None = None
) -> tuple[bytes, str]:
    matrix = await get_dashboard_provider().variance_discrepancy_subcategory_matrix(db, review_cycle_id=review_cycle_id)
    header, rows = _reason_matrix_grid(matrix)
    return _grid_to_xlsx("Discrepancy reasons", header, rows), f"discrepancy_reasons{_suffix(review_cycle_id)}_{date.today()}.xlsx"
