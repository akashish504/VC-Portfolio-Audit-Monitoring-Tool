"""Variance "Current-cycle status" breakdown export.

Builds an XLSX of the scoped-in cohort, one row per (deal, review status), so the
per-status deal counts reconcile exactly with the dashboard status breakdown and
its company drill-down (both count distinct deals, not entities).
"""
from __future__ import annotations

import io
from datetime import date
from typing import Optional

from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from src.db.models import Entity, File, PortfolioCompany


def _scoped_in_subquery(review_cycle_id: Optional[str]):
    """Companies with >=1 uploaded financials File for the cycle (the scoped-in cohort).

    Mirrors ``_scoped_in_clause`` in the dashboard provider so the export cohort
    matches the breakdown exactly.
    """
    sub = select(File.portfolio_company_id).where(File.portfolio_company_id.isnot(None))
    if review_cycle_id:
        sub = sub.where(File.review_cycle_id == review_cycle_id)
    return PortfolioCompany.id.in_(sub)


# (db column expression label, header) — order defines the sheet columns.
_COLUMNS: list[tuple[str, str]] = [
    ("status", "Review Status"),
    ("company_id", "CID"),
    ("name", "Deal Name"),
    ("fund", "Fund"),
    ("geography", "Geography"),
    ("fy_end", "FYE"),
    ("strategy", "Strategy"),
    ("investment_lead", "Investment Lead"),
    ("category", "Category"),
]


async def build_variance_status_breakdown_xlsx(
    db: AsyncSession, review_cycle_id: Optional[str]
) -> tuple[bytes, str]:
    """Return (xlsx_bytes, filename) for the variance review-status breakdown."""
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - openpyxl is a hard dep in prod
        raise RuntimeError("openpyxl is not installed") from exc

    # One distinct (status, deal) pair per row: grouping the sheet by Review Status
    # and counting distinct deals reproduces the dashboard breakdown numbers.
    stmt = (
        select(
            func.trim(Entity.status),
            PortfolioCompany.company_id,
            PortfolioCompany.name,
            PortfolioCompany.fund,
            PortfolioCompany.geography,
            PortfolioCompany.fy_end,
            PortfolioCompany.investment_stage,
            PortfolioCompany.investment_lead,
            PortfolioCompany.company_phase_category,
        )
        .join(PortfolioCompany, PortfolioCompany.id == Entity.portfolio_company_id)
        .where(_scoped_in_subquery(review_cycle_id), Entity.status.isnot(None))
        .distinct()
    )
    if review_cycle_id:
        stmt = stmt.where(PortfolioCompany.review_cycle_id == review_cycle_id)
    stmt = stmt.order_by(func.trim(Entity.status), PortfolioCompany.name)

    rows = (await db.execute(stmt)).all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Review Status"

    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="1D4ED8")
    hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    headers = [label for _field, label in _COLUMNS]
    for ci, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = hdr_align
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "A2"

    for r in rows:
        ws.append(["" if v is None else v for v in r])

    for ci, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(ci)].width = max(14, min(48, len(h) + 6))

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    suffix = f"_{review_cycle_id}" if review_cycle_id else ""
    filename = f"review_status_breakdown{suffix}_{date.today()}.xlsx"
    return buf.read(), filename
