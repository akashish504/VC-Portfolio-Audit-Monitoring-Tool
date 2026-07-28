from __future__ import annotations

import asyncio
import io
import logging
from datetime import date
from typing import Any, Optional, Tuple

from fastapi import APIRouter, BackgroundTasks, Body, Depends, File as FastAPIFile, Form, HTTPException, Query, UploadFile, status
from fastapi.responses import Response, StreamingResponse
from sqlalchemy import func, select
from sqlalchemy.ext.asyncio import AsyncSession

from src.db.session import get_db

logger = logging.getLogger(__name__)
from src.schema.common import Page
from src.schema.portfolio import (
    AuditFinancialAddCompositeFieldRequest,
    AuditFinancialDismissUnmatchedRequest,
    AuditFinancialExtractedValuePatchRequest,
    AuditFinancialFieldComponentsRequest,
    AuditFinancialFieldDetachRequest,
    AuditFinancialRestoreUnmatchedRequest,
    AuditFinancialUnmatchedMapRequest,
    BulkFileUploadResponse,
    ConvertCurrencyRequest,
    EntityReviewCycleConvertCurrencyRequest,
    EntityCreate,
    EntityPatch,
    EntityRead,
    ExtractionCurrencyPatchRequest,
    ExtractionConvertCurrencyApplyRequest,
    FileCreate,
    FilePatch,
    FileRead,
    FinancialDataSnowflakeCreate,
    FinancialDataSnowflakePatch,
    FinancialDataSnowflakeRead,
    FinancialMetricReconciliationPatch,
    FinancialMetricReconciliationRead,
    ManualReconciliationQueryCreate,
    ManualReconciliationQueryPatch,
    ManualReconciliationQueryRead,
    PortfolioCompanyBulkUpsertRequest,
    PortfolioCompanyBulkUpsertResponse,
    OrgChartClearResponse,
    PortfolioCompanyCreate,
    PortfolioCompanyPatch,
    PortfolioCompanyRead,
    CycleEntityRowRead,
    InReviewTrackerRowRead,
    SnowflakeEntityAttachmentPatch,
)
from src.db.file_filters import ORG_CHART_FILE_TAG, ORG_CHART_BATCH_FILE_TAG, exclude_org_chart_uploads_clause
from src.db.models import File as DbFile, OrgChartUploadBatch, OrgChartUploadRecord
from src.services.audit_service import (
    AuditService,
    MAX_BULK_AUDIT_UPLOAD_FILES,
    audit_upload_filename_key,
    schedule_extraction_with_background,
    validate_audit_upload_filename,
)
from src.utils.s3 import download_storage_uri
from src.services.portfolio import PortfolioService
from src.services.org_chart_extraction import run_org_chart_batch_record_extraction

router = APIRouter()


def _coerce_optional_positive_int(value: Any) -> Optional[int]:
    """Parse JSON numbers/strings into a positive entity/company id (or None)."""
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value if value > 0 else None
    if isinstance(value, float):
        if value != value or value in (float("inf"), float("-inf")):
            return None
        if not value.is_integer():
            return None
        i = int(value)
        return i if i > 0 else None
    if isinstance(value, str):
        s = value.strip()
        if not s:
            return None
        try:
            i = int(s, 10)
            return i if i > 0 else None
        except ValueError:
            return None
    return None


def _file_read_from_orm(f: DbFile) -> FileRead:
    base = FileRead.model_validate(f)
    pc = getattr(f, "portfolio_company", None)
    ent = getattr(f, "entity", None)
    # ocr_metadata (selectin-loaded) is written when extraction completes, so its updated_at is
    # the "last processed" time. None if the file has never been processed.
    ocr_meta = getattr(f, "ocr_metadata", None)
    return base.model_copy(
        update={
            "portfolio_company_name": pc.name if pc else None,
            "portfolio_company_review_cycle_id": pc.review_cycle_id if pc else None,
            "entity_name": ent.name if ent else None,
            "entity_geolocation": ent.geolocation if ent else None,
            "entity_fy_end": ent.fy_end if ent else None,
            "processed_at": ocr_meta.updated_at if ocr_meta else None,
        }
    )


def _pagination(
    limit: int = Query(50, ge=1, le=500),
    offset: int = Query(0, ge=0),
):
    return limit, offset


def _pagination_large(
    limit: int = Query(500, ge=1, le=5000),
    offset: int = Query(0, ge=0),
):
    """Higher limit cap for bulk-fetch endpoints (e.g. reconciliation rows)."""
    return limit, offset


@router.get("/portfolio-companies", response_model=Page[PortfolioCompanyRead])
async def list_portfolio_companies(
    q: Optional[str] = Query(default=None, description="Search by company_id or name"),
    review_cycle_id: Optional[str] = Query(default=None),
    review_stage: Optional[str] = Query(default=None),
    has_discrepancy_type: Optional[str] = Query(
        default=None,
        description="Only companies with at least one discrepancy of this type (e.g. revenue, ebitda)",
    ),
    has_discrepancy_category: Optional[str] = Query(
        default=None,
        description="Only companies with at least one discrepancy in this category (e.g. manual)",
    ),
    pagination: Tuple[int, int] = Depends(_pagination),
    db: AsyncSession = Depends(get_db),
):
    limit, offset = pagination
    items, total = await PortfolioService.list_portfolio_companies(
        db,
        q=q,
        review_cycle_id=review_cycle_id,
        review_stage=review_stage,
        has_discrepancy_type=has_discrepancy_type,
        has_discrepancy_category=has_discrepancy_category,
        limit=limit,
        offset=offset,
    )
    return Page(items=items, total=total)


@router.get("/cycle-entities", response_model=Page[CycleEntityRowRead])
async def list_cycle_entities(
    q: Optional[str] = Query(default=None, description="Search by company name/id or entity name"),
    review_cycle_id: Optional[str] = Query(default=None),
    status: Optional[str] = Query(default=None, description="Filter to entities with this exact status"),
    pagination: Tuple[int, int] = Depends(_pagination),
    db: AsyncSession = Depends(get_db),
):
    """Audit Tracker data source: every entity in the review cycle, with company context."""
    limit, offset = pagination
    items, total = await PortfolioService.list_cycle_entities(
        db,
        q=q,
        review_cycle_id=review_cycle_id,
        status=status,
        limit=limit,
        offset=offset,
    )
    return Page(items=items, total=total)


@router.get("/in-review-tracker", response_model=Page[InReviewTrackerRowRead])
async def list_in_review_tracker(
    q: Optional[str] = Query(default=None, description="Search by company name or company_id"),
    review_cycle_id: Optional[str] = Query(default=None),
    entity_status: Optional[str] = Query(default=None),
    pagination: Tuple[int, int] = Depends(_pagination),
    db: AsyncSession = Depends(get_db),
):
    limit, offset = pagination
    items, total = await PortfolioService.list_in_review_tracker_rows(
        db,
        q=q,
        review_cycle_id=review_cycle_id,
        entity_status=entity_status,
        limit=limit,
        offset=offset,
    )
    return Page(items=items, total=total)


@router.post(
    "/portfolio-companies",
    response_model=PortfolioCompanyRead,
    status_code=status.HTTP_201_CREATED,
)
async def create_portfolio_company(
    payload: PortfolioCompanyCreate,
    db: AsyncSession = Depends(get_db),
):
    return await PortfolioService.create_portfolio_company(db, payload)


@router.post(
    "/portfolio-companies/bulk-upsert",
    response_model=PortfolioCompanyBulkUpsertResponse,
    tags=["Portfolio"],
)
async def bulk_upsert_portfolio_companies(
    payload: PortfolioCompanyBulkUpsertRequest,
    db: AsyncSession = Depends(get_db),
):
    items = await PortfolioService.bulk_upsert_portfolio_companies(db, payload)
    return PortfolioCompanyBulkUpsertResponse(items=items)


async def _company_read_with_orphan_flag(db: AsyncSession, obj) -> PortfolioCompanyRead:
    orphan_count = (
        await db.execute(
            select(func.count(DbFile.id)).where(
                DbFile.portfolio_company_id == obj.id,
                DbFile.entity_id.is_(None),
                (DbFile.entity_detached_acknowledged.is_(None) | DbFile.entity_detached_acknowledged.is_(False)),
                exclude_org_chart_uploads_clause(),
            )
        )
    ).scalar_one()
    result = PortfolioCompanyRead.model_validate(obj)
    result.has_orphan_files = orphan_count > 0
    return result


@router.get("/portfolio-companies/{portfolio_company_id}", response_model=PortfolioCompanyRead)
async def get_portfolio_company(
    portfolio_company_id: int,
    db: AsyncSession = Depends(get_db),
):
    obj = await PortfolioService.get_portfolio_company(db, portfolio_company_id)
    if not obj:
        raise HTTPException(status_code=404, detail="PortfolioCompany not found")
    return await _company_read_with_orphan_flag(db, obj)


@router.get(
    "/portfolio-companies/by-company-id/{business_company_id}",
    response_model=PortfolioCompanyRead,
)
async def get_portfolio_company_by_business_id(
    business_company_id: str,
    review_cycle_id: Optional[str] = Query(default=None),
    db: AsyncSession = Depends(get_db),
):
    obj = await PortfolioService.get_portfolio_company_by_company_id(
        db,
        business_company_id,
        review_cycle_id=review_cycle_id,
    )
    if not obj:
        raise HTTPException(status_code=404, detail="PortfolioCompany not found")
    return await _company_read_with_orphan_flag(db, obj)


@router.patch("/portfolio-companies/{portfolio_company_id}", response_model=PortfolioCompanyRead)
async def patch_portfolio_company(
    portfolio_company_id: int,
    payload: PortfolioCompanyPatch,
    db: AsyncSession = Depends(get_db),
):
    obj = await PortfolioService.patch_portfolio_company(db, portfolio_company_id, payload)
    if not obj:
        raise HTTPException(status_code=404, detail="PortfolioCompany not found")
    return obj


@router.delete("/portfolio-companies/{portfolio_company_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_portfolio_company(
    portfolio_company_id: int,
    db: AsyncSession = Depends(get_db),
):
    ok = await PortfolioService.delete_portfolio_company(db, portfolio_company_id)
    if not ok:
        raise HTTPException(status_code=404, detail="PortfolioCompany not found")
    return None


@router.get("/portfolio-companies/{portfolio_company_id}/org-entities", response_model=list[EntityRead])
async def list_org_entities(
    portfolio_company_id: int,
    db: AsyncSession = Depends(get_db),
):
    # Compatibility endpoint used by org chart UI.
    return await PortfolioService.list_entities_for_company(db, portfolio_company_id)


@router.post(
    "/portfolio-companies/{portfolio_company_id}/org-entities/reparent",
    response_model=list[EntityRead],
)
async def reparent_org_entity(
    portfolio_company_id: int,
    payload: dict,
    db: AsyncSession = Depends(get_db),
):
    child_entity_id = payload.get("child_entity_id")
    new_parent_entity_id = payload.get("new_parent_entity_id")
    if not isinstance(child_entity_id, int):
        raise HTTPException(status_code=422, detail="child_entity_id must be an integer")
    if new_parent_entity_id is not None and not isinstance(new_parent_entity_id, int):
        raise HTTPException(status_code=422, detail="new_parent_entity_id must be an integer or null")
    items = await PortfolioService.reparent_entity(
        db,
        portfolio_company_id=portfolio_company_id,
        child_entity_id=child_entity_id,
        new_parent_entity_id=new_parent_entity_id,
    )
    if items is None:
        raise HTTPException(status_code=404, detail="Entity not found")
    return items


@router.post(
    "/portfolio-companies/{portfolio_company_id}/org-chart/upload",
    tags=["Org chart"],
)
async def upload_org_chart(
    portfolio_company_id: int,
    background_tasks: BackgroundTasks,
    db: AsyncSession = Depends(get_db),
    upload: UploadFile = FastAPIFile(...),
):
    file_bytes = await upload.read()
    filename = upload.filename or "org-chart"
    svc = AuditService(db)

    # Upload the file with the batch-staging tag so it is not shown in the
    # company file list until reconciliation promotes it to ORG_CHART_FILE_TAG.
    row = await svc.upload_file_direct(
        company_id=portfolio_company_id,
        file_name=filename,
        file_bytes=file_bytes,
        mime_type=upload.content_type,
        extra_tags=[ORG_CHART_BATCH_FILE_TAG],
    )

    # Create a sentinel batch + record so the reconciliation flow (which is
    # driven by OrgChartUploadRecord) works identically to ZIP batch uploads.
    batch = OrgChartUploadBatch(
        review_cycle_id=None,
        original_zip_filename="__manual_upload__",
        status="uploaded",
        file_count=1,
    )
    db.add(batch)
    await db.flush()
    await db.refresh(batch)

    record = OrgChartUploadRecord(
        batch_id=batch.id,
        review_cycle_id=None,
        file_id=row.id,
        file_name=filename,
        storage_uri=row.storage_uri,
        company_id=str(portfolio_company_id),
        name=filename,
        portfolio_company_id=portfolio_company_id,
        extracted_org_chart=None,
        extraction_status="mapped",
        reconciliation_status=None,
        error_message=None,
    )
    db.add(record)
    await db.flush()
    await db.refresh(record)
    record_id = record.id
    await db.commit()

    # Trigger extraction via the batch record path — on completion this stores
    # extracted entities in the record and triggers the reconciliation dialog.
    background_tasks.add_task(_run_manual_upload_extraction, record_id)

    dl = await svc.get_download_url(row.id, expires_in=3600)
    return {
        "file_id": row.id,
        "record_id": record_id,
        "status": "queued",
        "updated_at": (row.updated_at or row.created_at).isoformat() if (row.updated_at or row.created_at) else None,
        "filename": row.filename,
        "content_type": row.content_type,
        "download_url": dl.get("download_url"),
        "download_expires_in": dl.get("expires_in"),
    }


async def _run_manual_upload_extraction(record_id: int) -> None:
    from src.db.session import AsyncSessionLocal
    try:
        async with AsyncSessionLocal() as db:
            await run_org_chart_batch_record_extraction(db, record_id)
    except Exception:
        logger.exception("manual upload extraction failed record_id=%s", record_id)


@router.delete(
    "/portfolio-companies/{portfolio_company_id}/org-chart",
    tags=["Org chart"],
    response_model=OrgChartClearResponse,
    summary="Remove org chart file and delete all entities for this company",
)
async def clear_org_chart(
    portfolio_company_id: int,
    db: AsyncSession = Depends(get_db),
):
    out = await PortfolioService.clear_org_chart_and_entities(db, portfolio_company_id=portfolio_company_id)
    if out is None:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail="Portfolio company not found")
    return OrgChartClearResponse(
        deleted_entities=int(out["deleted_entities"]),
        deleted_file_id=out["deleted_file_id"],
    )


@router.post(
    "/portfolio-companies/{portfolio_company_id}/org-chart/retrigger",
    tags=["Org chart"],
)
async def retrigger_org_chart(
    portfolio_company_id: int,
    background_tasks: BackgroundTasks,
    file_id: Optional[int] = Query(default=None),
    db: AsyncSession = Depends(get_db),
):
    svc = AuditService(db)
    if file_id is None:
        raise HTTPException(status_code=422, detail="file_id query param is required")
    row = await svc._get_file_or_raise(file_id)
    if row.portfolio_company_id != portfolio_company_id:
        raise HTTPException(status_code=404, detail="File not found")
    await schedule_extraction_with_background(
        audit_service=svc,
        background_tasks=background_tasks,
        file_id=row.id,
        kind="org_chart",
    )
    dl = await svc.get_download_url(row.id, expires_in=3600)
    return {
        "file_id": row.id,
        "status": "queued",
        "updated_at": (row.updated_at or row.created_at).isoformat() if (row.updated_at or row.created_at) else None,
        "filename": row.filename,
        "content_type": row.content_type,
        "download_url": dl.get("download_url"),
        "download_expires_in": dl.get("expires_in"),
    }


@router.post(
    "/files/init-upload",
    tags=["Files"],
    summary="Register pending upload; send bytes only via POST /api/v1/files/{file_id}/upload",
)
async def init_upload(payload: dict, db: AsyncSession = Depends(get_db)):
    filename = payload.get("filename")
    content_type = payload.get("content_type")
    portfolio_company_id = payload.get("portfolio_company_id")
    review_cycle_id = payload.get("review_cycle_id")
    fy_end = payload.get("fy_end")
    raw_entity = payload.get("entity_id")
    entity_id = _coerce_optional_positive_int(raw_entity)

    if not isinstance(filename, str) or not filename.strip():
        raise HTTPException(status_code=422, detail="filename is required")

    if portfolio_company_id is not None and not isinstance(portfolio_company_id, int):
        raise HTTPException(status_code=422, detail="portfolio_company_id must be an integer or null")

    company_id: Optional[int] = portfolio_company_id if isinstance(portfolio_company_id, int) else None

    if entity_id is not None and company_id is None:
        raise HTTPException(
            status_code=422,
            detail="entity_id requires portfolio_company_id",
        )

    svc = AuditService(db)
    try:
        out = await svc.generate_upload_url(
            company_id=company_id,
            file_name=filename,
            mime_type=content_type,
            entity_id=entity_id,
            review_cycle_id=review_cycle_id if isinstance(review_cycle_id, str) and review_cycle_id.strip() else None,
            fy_end=fy_end if isinstance(fy_end, str) and fy_end.strip() else None,
        )
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e)) from None

    return out


@router.post(
    "/files/bulk-upload",
    response_model=BulkFileUploadResponse,
    tags=["Files"],
    summary="Upload up to 10 audit files in one request (auto-starts extraction per success)",
)
async def bulk_upload_audit_files(
    background_tasks: BackgroundTasks,
    fy_end: Optional[str] = Form(None),
    review_cycle_id: Optional[str] = Form(None),
    kind: str = Form("audit_report"),
    portfolio_company_id: Optional[int] = Form(None),
    entity_id: Optional[int] = Form(None),
    uploads: list[UploadFile] = FastAPIFile(...),
    db: AsyncSession = Depends(get_db),
):
    del kind  # reserved; extraction always uses audit_financials
    from src.services.fy_end import normalize_fy_end, resolve_review_cycle_id_for_fy_end

    coerced_entity = _coerce_optional_positive_int(entity_id)
    coerced_company = _coerce_optional_positive_int(portfolio_company_id)

    fe = (fy_end or "").strip()
    normalized_fy: Optional[str] = None
    if fe:
        normalized_fy = normalize_fy_end(fe)
        if not normalized_fy:
            raise HTTPException(status_code=422, detail=f"invalid fy_end: {fe!r}")
    elif coerced_company is not None:
        raise HTTPException(status_code=422, detail="fy_end is required when attaching to a company")

    rc = (review_cycle_id or "").strip()
    if not rc:
        if not normalized_fy:
            raise HTTPException(status_code=422, detail="review_cycle_id is required when fy_end is not provided")
        try:
            rc = await resolve_review_cycle_id_for_fy_end(db, normalized_fy)
        except ValueError as e:
            raise HTTPException(status_code=422, detail=str(e)) from None

    if coerced_entity is not None and coerced_company is None:
        raise HTTPException(
            status_code=422,
            detail="entity_id requires portfolio_company_id",
        )

    if not uploads:
        raise HTTPException(status_code=422, detail="At least one file is required")
    if len(uploads) > MAX_BULK_AUDIT_UPLOAD_FILES:
        raise HTTPException(
            status_code=422,
            detail=f"At most {MAX_BULK_AUDIT_UPLOAD_FILES} files per bulk upload",
        )

    seen_names: set[str] = set()
    duplicate_names: list[str] = []
    for up in uploads:
        raw_name = up.filename or "upload"
        key = audit_upload_filename_key(raw_name)
        if key in seen_names:
            duplicate_names.append(raw_name)
        seen_names.add(key)
    if duplicate_names:
        raise HTTPException(
            status_code=422,
            detail=f"Duplicate filenames in batch: {', '.join(sorted(set(duplicate_names)))}",
        )

    for up in uploads:
        err = validate_audit_upload_filename(up.filename or "")
        if err:
            raise HTTPException(status_code=422, detail=f"{up.filename or 'upload'}: {err}")

    if coerced_company is not None:
        try:
            await PortfolioService.update_company_fy_end(
                db,
                portfolio_company_id=coerced_company,
                fy_end=normalized_fy,
            )
        except ValueError as e:
            raise HTTPException(status_code=422, detail=str(e)) from None

    file_payloads: list[tuple[str, bytes, Optional[str]]] = []
    for up in uploads:
        data = await up.read()
        file_payloads.append((up.filename or "upload", data, up.content_type))

    svc = AuditService(db)
    try:
        out = await svc.bulk_upload_audit_files(
            company_id=coerced_company,
            entity_id=coerced_entity,
            files=file_payloads,
            review_cycle_id=rc or None,
            fy_end=normalized_fy,
        )
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e)) from None

    for item in out.get("results") or []:
        if item.get("status") == "success" and item.get("file_id"):
            await schedule_extraction_with_background(
                audit_service=svc,
                background_tasks=background_tasks,
                file_id=int(item["file_id"]),
                kind="audit_financials",
            )

    return out


@router.post("/files/{file_id}/upload", tags=["Files"])
async def upload_existing_file(
    file_id: int,
    upload: UploadFile = FastAPIFile(...),
    db: AsyncSession = Depends(get_db),
):
    file_bytes = await upload.read()
    svc = AuditService(db)
    row = await svc.upload_to_existing(file_id=file_id, file_bytes=file_bytes, content_type=upload.content_type)
    return {"file_id": row.id, "status": row.status}


@router.post("/files/{file_id}/extract", tags=["Files"])
async def start_extraction(
    file_id: int,
    payload: dict,
    background_tasks: BackgroundTasks,
    db: AsyncSession = Depends(get_db),
):
    kind = payload.get("kind")
    if not isinstance(kind, str) or not kind.strip():
        raise HTTPException(status_code=422, detail="kind is required")
    svc = AuditService(db)
    return await schedule_extraction_with_background(
        audit_service=svc,
        background_tasks=background_tasks,
        file_id=file_id,
        kind=kind,
    )


@router.get("/files/{file_id}/extract/status", tags=["Files"])
async def extraction_status(file_id: int, db: AsyncSession = Depends(get_db)):
    return await AuditService(db).extraction_status(file_id=file_id)


@router.patch(
    "/files/{file_id}/extraction/currency",
    tags=["Files"],
    summary="Manually set the ISO-4217 currency label on a file's extraction metadata",
)
async def set_extraction_currency(
    file_id: int,
    payload: ExtractionCurrencyPatchRequest,
    db: AsyncSession = Depends(get_db),
):
    svc = AuditService(db)
    try:
        return await svc.set_extraction_currency(file_id=file_id, currency=payload.currency)
    except FileNotFoundError as e:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(e))
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))


@router.post(
    "/files/{file_id}/extraction/convert-currency",
    tags=["Files"],
    summary="Apply FX conversion to extracted P&L / balance sheet / cash flow amounts (persists)",
)
async def apply_extraction_currency_conversion(
    file_id: int,
    payload: ExtractionConvertCurrencyApplyRequest,
    db: AsyncSession = Depends(get_db),
):
    svc = AuditService(db)
    try:
        return await svc.apply_extraction_currency_conversion(file_id=file_id, target_currency=payload.target_currency)
    except FileNotFoundError as e:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(e))
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))


@router.patch(
    "/files/{file_id}/audit-financials/extracted-value",
    tags=["Files"],
    summary="Set or clear one numeric field in the extracted financials tree (manual correction)",
)
async def patch_audit_financial_extracted_value(
    file_id: int,
    payload: AuditFinancialExtractedValuePatchRequest,
    db: AsyncSession = Depends(get_db),
):
    svc = AuditService(db)
    try:
        return await svc.patch_audit_financial_extracted_value(
            file_id=file_id,
            path=payload.path.strip(),
            value=payload.value,
            reason=payload.reason,
        )
    except FileNotFoundError as e:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(e))
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))


@router.post(
    "/files/{file_id}/audit-financials/add-composite-field",
    tags=["Files"],
    summary="Create an empty composite (bucket) node in the extracted financials tree",
)
async def audit_financial_add_composite_field(
    file_id: int,
    payload: AuditFinancialAddCompositeFieldRequest,
    db: AsyncSession = Depends(get_db),
):
    svc = AuditService(db)
    try:
        return await svc.add_audit_financial_composite_field(
            file_id=file_id,
            path=payload.path.strip(),
            reason=payload.reason,
        )
    except FileNotFoundError as e:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(e))
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))


@router.get(
    "/files/{file_id}/audit-financials/parent-paths",
    tags=["Files"],
    summary="List canonical dict paths for mapping unmatched financial lines",
)
async def audit_financial_parent_paths(file_id: int, db: AsyncSession = Depends(get_db)):
    svc = AuditService(db)
    await svc._get_file_or_raise(file_id)
    result = await svc.get_audit_financial_parent_paths(file_id=file_id)
    return result


@router.get(
    "/files/{file_id}/source-refs",
    tags=["Files"],
    summary="Return the stored source-ref map for a file (page + text snippet per extracted field path)",
)
async def file_source_refs(file_id: int, db: AsyncSession = Depends(get_db)):
    """
    Returns ``{file_id, source_refs, status}`` where ``source_refs`` is a flat dict mapping
    dotted field paths (e.g. ``profit_and_loss.revenue.revenue_from_operations``) to
    ``{page: int, text_snippet: str}``.

    ``status`` is one of: ``"completed"`` | ``"empty"`` | ``"skipped"`` | ``"not_available"``.
    """
    svc = AuditService(db)
    try:
        result = await svc.get_source_refs(file_id)
        return result
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail="File not found") from None


@router.get(
    "/entities/{entity_id}/afs-source",
    tags=["Files"],
    summary="Resolve an entity's audited-financials source file + its source-ref map",
)
async def entity_afs_source(entity_id: int, db: AsyncSession = Depends(get_db)):
    """
    Returns ``{entity_id, file_id, source_refs, status}`` for the entity's most-recent
    audited-financials file. ``source_refs`` maps dotted field paths (e.g.
    ``balance_sheet.assets.current_assets.financial_assets.cash_and_cash_equivalents``) to
    ``{page: int, text_snippet: str}`` so the UI can open the PDF source viewer for a path.

    ``file_id`` is ``null`` when the entity has no audited-financials file (UI keeps paths
    non-clickable). ``status`` is one of: ``"completed"`` | ``"empty"`` | ``"skipped"`` |
    ``"not_available"``. ``files`` lists the entity's candidate AFS files (newest first) with the
    current primary flagged, for the dashboard's file picker.
    """
    return await AuditService(db).get_entity_afs_source(entity_id)


@router.put(
    "/entities/{entity_id}/afs-primary-file",
    tags=["Files"],
    summary="Choose which AFS file is the reconciliation source for an entity (drives values + emails)",
)
async def set_entity_primary_afs_file(
    entity_id: int,
    file_id: int = Body(..., embed=True),
    db: AsyncSession = Depends(get_db),
):
    """Make ``file_id`` the authoritative audited-financials file for its (entity, cycle) and
    re-sync, so the dashboard values, breakdowns, source links and query-email numbers all come
    from the chosen file. Returns the refreshed ``get_entity_afs_source`` payload.
    """
    try:
        return await AuditService(db).set_entity_primary_afs_file(entity_id=entity_id, file_id=file_id)
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail="File not found") from None
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc


@router.post(
    "/files/{file_id}/audit-financials/map-unmatched",
    tags=["Files"],
    summary="Place an unmatched line under a canonical parent path (human-in-the-loop)",
)
async def audit_financial_map_unmatched(
    file_id: int,
    payload: AuditFinancialUnmatchedMapRequest,
    db: AsyncSession = Depends(get_db),
):
    svc = AuditService(db)
    try:
        out = await svc.map_audit_financial_unmatched(
            file_id,
            unmatched_id=payload.unmatched_id,
            target_parent_path=payload.target_parent_path,
            target_key=payload.target_key,
            confirm_overwrite=payload.confirm_overwrite,
            conflict_mode=payload.conflict_mode,
            reason=payload.reason,
        )
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))
    if out.get("conflict"):
        raise HTTPException(status_code=status.HTTP_409_CONFLICT, detail=out)
    return out


@router.post(
    "/files/{file_id}/audit-financials/detach-field",
    tags=["Files"],
    summary="Remove a field from the extracted tree and move it to the unmatched panel (recoverable)",
)
async def audit_financial_detach_field(
    file_id: int,
    payload: AuditFinancialFieldDetachRequest,
    db: AsyncSession = Depends(get_db),
):
    svc = AuditService(db)
    try:
        return await svc.detach_audit_financial_field(file_id, path=payload.path)
    except FileNotFoundError as e:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(e))
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))


@router.post(
    "/files/{file_id}/audit-financials/dismiss-unmatched",
    tags=["Files"],
    summary="Dismiss an unmatched row (moves to dismissed section, always recoverable)",
)
async def audit_financial_dismiss_unmatched(
    file_id: int,
    payload: AuditFinancialDismissUnmatchedRequest,
    db: AsyncSession = Depends(get_db),
):
    svc = AuditService(db)
    try:
        return await svc.dismiss_audit_financial_unmatched(file_id, unmatched_id=payload.unmatched_id)
    except FileNotFoundError as e:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(e))
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))


@router.post(
    "/files/{file_id}/audit-financials/restore-unmatched",
    tags=["Files"],
    summary="Restore a dismissed unmatched row back to the active unmatched list",
)
async def audit_financial_restore_unmatched(
    file_id: int,
    payload: AuditFinancialRestoreUnmatchedRequest,
    db: AsyncSession = Depends(get_db),
):
    svc = AuditService(db)
    try:
        return await svc.restore_audit_financial_unmatched(file_id, unmatched_id=payload.unmatched_id)
    except FileNotFoundError as e:
        raise HTTPException(status_code=status.HTTP_404_NOT_FOUND, detail=str(e))
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))


@router.post(
    "/files/{file_id}/audit-financials/field-components",
    tags=["Files"],
    summary="Replace a canonical field's roll-up breakdown (leaf becomes the signed sum)",
)
async def audit_financial_set_field_components(
    file_id: int,
    payload: AuditFinancialFieldComponentsRequest,
    db: AsyncSession = Depends(get_db),
):
    svc = AuditService(db)
    try:
        return await svc.set_audit_financial_field_components(
            file_id,
            path=payload.path,
            components=[c.model_dump() for c in payload.components],
            reason=payload.reason,
        )
    except ValueError as e:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(e))


@router.get("/entities", response_model=Page[EntityRead])
async def list_entities(
    portfolio_company_id: Optional[int] = Query(default=None),
    review_cycle: Optional[str] = Query(default=None),
    pagination: Tuple[int, int] = Depends(_pagination),
    db: AsyncSession = Depends(get_db),
):
    limit, offset = pagination
    items, total = await PortfolioService.list_entities(
        db,
        portfolio_company_id=portfolio_company_id,
        review_cycle=review_cycle,
        limit=limit,
        offset=offset,
    )
    reads = []
    for ent in items:
        r = EntityRead.model_validate(ent)
        pc = getattr(ent, "portfolio_company", None)
        r = r.model_copy(update={"portfolio_company_name": pc.name if pc else None})
        reads.append(r)
    return Page(items=reads, total=total)


@router.post("/entities", response_model=EntityRead, status_code=status.HTTP_201_CREATED)
async def create_entity(
    payload: EntityCreate,
    db: AsyncSession = Depends(get_db),
):
    obj = await PortfolioService.create_entity(db, payload)
    if not obj:
        raise HTTPException(status_code=404, detail="PortfolioCompany not found")
    return obj


@router.get("/entities/{entity_id}", response_model=EntityRead)
async def get_entity(
    entity_id: int,
    db: AsyncSession = Depends(get_db),
):
    obj = await PortfolioService.get_entity(db, entity_id)
    if not obj:
        raise HTTPException(status_code=404, detail="Entity not found")
    return obj


@router.patch("/entities/{entity_id}", response_model=EntityRead)
async def patch_entity(
    entity_id: int,
    payload: EntityPatch,
    db: AsyncSession = Depends(get_db),
):
    obj = await PortfolioService.patch_entity(db, entity_id, payload)
    if not obj:
        raise HTTPException(status_code=404, detail="Entity not found")
    return obj


@router.delete("/entities/{entity_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_entity(
    entity_id: int,
    db: AsyncSession = Depends(get_db),
):
    ok = await PortfolioService.delete_entity(db, entity_id)
    if not ok:
        raise HTTPException(status_code=404, detail="Entity not found")
    return None


# ---- Financial Metric Reconciliation ----

@router.get("/financial-metric-reconciliation", response_model=Page, tags=["Reconciliation"])
async def list_financial_metric_reconciliation(
    portfolio_company_id: Optional[int] = Query(default=None),
    entity_id: Optional[int] = Query(default=None),
    pagination: Tuple[int, int] = Depends(_pagination_large),
    db: AsyncSession = Depends(get_db),
):
    limit, offset = pagination
    rows, total = await PortfolioService.list_financial_metric_reconciliation(
        db,
        portfolio_company_id=portfolio_company_id,
        entity_id=entity_id,
        limit=limit,
        offset=offset,
    )
    entity_names: dict[int, str] = {}
    for r in rows:
        if r.entity and r.entity_id not in entity_names:
            entity_names[r.entity_id] = r.entity.name
    items = await PortfolioService.build_reconciliation_reads(
        db, rows, entity_names=entity_names
    )
    return Page(items=items, total=total)


@router.get("/financial-metric-reconciliation/{reconciliation_id}", response_model=None, tags=["Reconciliation"])
async def get_financial_metric_reconciliation(
    reconciliation_id: int,
    db: AsyncSession = Depends(get_db),
):
    obj = await PortfolioService.get_financial_metric_reconciliation(db, reconciliation_id)
    if not obj:
        raise HTTPException(status_code=404, detail="Reconciliation row not found")
    entity_names = {obj.entity_id: obj.entity.name} if obj.entity else {}
    items = await PortfolioService.build_reconciliation_reads(db, [obj], entity_names=entity_names)
    return items[0]


@router.patch("/financial-metric-reconciliation/{reconciliation_id}", response_model=None, tags=["Reconciliation"])
async def patch_financial_metric_reconciliation(
    reconciliation_id: int,
    payload: FinancialMetricReconciliationPatch,
    db: AsyncSession = Depends(get_db),
):
    try:
        obj = await PortfolioService.patch_financial_metric_reconciliation(
            db, reconciliation_id, payload
        )
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    if not obj:
        raise HTTPException(status_code=404, detail="Reconciliation row not found")
    await db.commit()
    entity_names = {obj.entity_id: obj.entity.name} if obj.entity else {}
    items = await PortfolioService.build_reconciliation_reads(db, [obj], entity_names=entity_names)
    return items[0]


@router.delete("/financial-metric-reconciliation/{reconciliation_id}", status_code=status.HTTP_204_NO_CONTENT, tags=["Reconciliation"])
async def delete_financial_metric_reconciliation(
    reconciliation_id: int,
    db: AsyncSession = Depends(get_db),
):
    ok = await PortfolioService.delete_financial_metric_reconciliation(db, reconciliation_id)
    if not ok:
        raise HTTPException(status_code=404, detail="Reconciliation row not found")
    await db.commit()
    return None


@router.post(
    "/financial-metric-reconciliation/convert-afs-currency",
    response_model=None,
    tags=["Reconciliation"],
)
async def convert_afs_currency_for_entity_cycle(
    payload: EntityReviewCycleConvertCurrencyRequest,
    db: AsyncSession = Depends(get_db),
):
    """Convert all AFS amounts for a given entity + review cycle to a new currency."""
    from src.db.models import Entity

    entity = await db.get(Entity, payload.entity_id)
    if entity is None:
        raise HTTPException(status_code=404, detail="Entity not found")
    try:
        updated = await PortfolioService.convert_afs_currency_for_entity_cycle(
            db,
            portfolio_company_id=entity.portfolio_company_id,
            entity_id=payload.entity_id,
            review_cycle=payload.review_cycle,
            target_currency=payload.target_currency,
        )
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    await db.commit()
    return {"updated": updated or 0}


# ---- Manual Reconciliation Queries ----

@router.get("/manual-reconciliation-queries", response_model=Page, tags=["Reconciliation"])
async def list_manual_reconciliation_queries(
    portfolio_company_id: Optional[int] = Query(default=None),
    entity_id: Optional[int] = Query(default=None),
    pagination: Tuple[int, int] = Depends(_pagination_large),
    db: AsyncSession = Depends(get_db),
):
    limit, offset = pagination
    items, total = await PortfolioService.list_manual_reconciliation_queries(
        db,
        portfolio_company_id=portfolio_company_id,
        entity_id=entity_id,
        limit=limit,
        offset=offset,
    )
    results = []
    for obj in items:
        d = ManualReconciliationQueryRead.model_validate(obj)
        if obj.entity:
            d.entity_name = obj.entity.name
        results.append(d)
    return Page(items=results, total=total)


@router.post("/manual-reconciliation-queries", response_model=None, status_code=status.HTTP_201_CREATED, tags=["Reconciliation"])
async def create_manual_reconciliation_query(
    payload: ManualReconciliationQueryCreate,
    db: AsyncSession = Depends(get_db),
):
    obj = await PortfolioService.create_manual_reconciliation_query(db, payload)
    if not obj:
        raise HTTPException(status_code=404, detail="PortfolioCompany not found")
    await db.commit()
    obj = await PortfolioService.get_manual_reconciliation_query(db, obj.id)
    d = ManualReconciliationQueryRead.model_validate(obj)
    if obj and obj.entity:
        d.entity_name = obj.entity.name
    return d


@router.patch("/manual-reconciliation-queries/{query_id}", response_model=None, tags=["Reconciliation"])
async def patch_manual_reconciliation_query(
    query_id: int,
    payload: "ManualReconciliationQueryPatch",
    db: AsyncSession = Depends(get_db),
):
    obj = await PortfolioService.patch_manual_reconciliation_query(db, query_id, payload)
    if not obj:
        raise HTTPException(status_code=404, detail="Manual query not found")
    await db.commit()
    await db.refresh(obj)
    d = ManualReconciliationQueryRead.model_validate(obj)
    if obj.entity:
        d.entity_name = obj.entity.name
    return d


@router.delete("/manual-reconciliation-queries/{query_id}", status_code=status.HTTP_204_NO_CONTENT, tags=["Reconciliation"])
async def delete_manual_reconciliation_query(
    query_id: int,
    db: AsyncSession = Depends(get_db),
):
    ok = await PortfolioService.delete_manual_reconciliation_query(db, query_id)
    if not ok:
        raise HTTPException(status_code=404, detail="Manual query not found")
    await db.commit()
    return None


# Snowflake financials kept (MIS source, entity attachment, currency conversion)

@router.post("/financial-data-snowflake", response_model=FinancialDataSnowflakeRead, status_code=status.HTTP_201_CREATED)
async def create_financial_data_snowflake(
    body: FinancialDataSnowflakeCreate,
    db: AsyncSession = Depends(get_db),
):
    result = await PortfolioService.create_financial_data_snowflake(
        db,
        portfolio_company_id=body.portfolio_company_id,
        entity_id=body.entity_id,
        review_cycle=body.review_cycle,
        frequency=body.frequency,
        currency=body.currency,
        metric_key=body.metric_key,
        metric_value=body.metric_value,
        edit_reason=body.edit_reason,
    )
    await db.commit()
    return result


@router.get("/financial-data-snowflake")
async def list_financial_data_snowflake(
    portfolio_company_id: Optional[int] = Query(default=None),
    entity_id: Optional[int] = Query(default=None),
    pagination: Tuple[int, int] = Depends(_pagination),
    db: AsyncSession = Depends(get_db),
):
    limit, offset = pagination
    items, total = await PortfolioService.list_financial_data_snowflake(
        db,
        portfolio_company_id=portfolio_company_id,
        entity_id=entity_id,
        limit=limit,
        offset=offset,
    )
    return {"items": items, "total": total}


@router.patch("/financial-data-snowflake/{row_id}", response_model=FinancialDataSnowflakeRead)
async def patch_financial_data_snowflake(
    row_id: int,
    body: FinancialDataSnowflakePatch,
    db: AsyncSession = Depends(get_db),
):
    patch_dict = body.model_dump(exclude_unset=True, exclude={"edit_reason"})
    result = await PortfolioService.patch_financial_data_snowflake(
        db,
        row_id=row_id,
        payload_patch=patch_dict,
        audit_financial_metric_edits=True,
        metric_edit_reason=body.edit_reason,
    )
    if result is None:
        raise HTTPException(status_code=404, detail="FinancialDataSnowflake row not found")
    await db.commit()
    return result


@router.patch("/financial-data-snowflake/{row_id}/entity-attachment", response_model=FinancialDataSnowflakeRead)
async def patch_snowflake_entity_attachment(
    row_id: int,
    body: SnowflakeEntityAttachmentPatch,
    db: AsyncSession = Depends(get_db),
):
    result = await PortfolioService.patch_financial_data_snowflake(
        db,
        row_id=row_id,
        payload_patch={"entity_id": body.entity_id},
        audit_financial_metric_edits=False,
    )
    if result is None:
        raise HTTPException(status_code=404, detail="FinancialDataSnowflake row not found")
    await db.commit()
    return result


@router.post("/financial-data-snowflake/{row_id}/convert-currency", response_model=FinancialDataSnowflakeRead)
async def convert_snowflake_currency(
    row_id: int,
    body: ConvertCurrencyRequest,
    db: AsyncSession = Depends(get_db),
):
    try:
        result = await PortfolioService.convert_financial_data_snowflake_currency(
            db,
            row_id=row_id,
            target_currency=body.target_currency,
        )
    except HTTPException:
        raise
    except ValueError as exc:
        raise HTTPException(status_code=422, detail=str(exc)) from exc
    if result is None:
        raise HTTPException(status_code=404, detail="FinancialDataSnowflake row not found")
    await db.commit()
    return result


@router.get("/files", response_model=Page[FileRead])
async def list_files(
    portfolio_company_id: Optional[int] = Query(default=None),
    entity_id: Optional[int] = Query(default=None),
    status_: Optional[str] = Query(default=None, alias="status"),
    unattached_only: bool = Query(default=False),
    pagination: Tuple[int, int] = Depends(_pagination),
    db: AsyncSession = Depends(get_db),
):
    limit, offset = pagination
    items, total = await PortfolioService.list_files(
        db,
        portfolio_company_id=portfolio_company_id,
        entity_id=entity_id,
        status_=status_,
        unattached_only=unattached_only,
        limit=limit,
        offset=offset,
    )
    return Page(items=[_file_read_from_orm(f) for f in items], total=total)


@router.post("/files", response_model=FileRead, status_code=status.HTTP_201_CREATED)
async def create_file(
    payload: FileCreate,
    db: AsyncSession = Depends(get_db),
):
    obj = await PortfolioService.create_file(db, payload)
    if not obj:
        raise HTTPException(status_code=404, detail="PortfolioCompany not found")
    reloaded = await PortfolioService.get_file(db, obj.id)
    return _file_read_from_orm(reloaded) if reloaded else _file_read_from_orm(obj)


@router.get("/files/{file_id}/download-url", tags=["Files"])
async def file_download_url(
    file_id: int,
    expires_in: int = Query(default=3600, ge=60, le=86400),
    db: AsyncSession = Depends(get_db),
):
    """Presigned GET URL for the file in S3 (same backing store as audit-files download)."""
    try:
        return await AuditService(db).get_download_url(file_id, expires_in=expires_in)
    except FileNotFoundError:
        raise HTTPException(status_code=404, detail="File not found") from None
    except ValueError as e:
        raise HTTPException(status_code=409, detail=str(e)) from e


@router.get("/files/{file_id}/stream", tags=["Files"])
async def stream_file(
    file_id: int,
    db: AsyncSession = Depends(get_db),
):
    """Stream the raw file bytes from S3 through the backend.

    Avoids browser CORS restrictions that block direct S3 presigned URL fetches
    from pdfjs / react-pdf when the S3 bucket has no AllowedOrigins CORS policy.
    """
    obj = await PortfolioService.get_file(db, file_id)
    if not obj:
        raise HTTPException(status_code=404, detail="File not found")
    if not obj.storage_uri:
        raise HTTPException(status_code=409, detail="File has no storage URI")
    try:
        data = await asyncio.to_thread(download_storage_uri, obj.storage_uri)
    except Exception as exc:
        raise HTTPException(status_code=502, detail="Could not fetch file from storage") from exc
    content_type = obj.content_type or "application/pdf"
    return Response(
        content=data,
        media_type=content_type,
        headers={"Content-Disposition": f'inline; filename="{obj.filename}"'},
    )


@router.get(
    "/files/{file_id}/preview/spreadsheet",
    tags=["Files"],
    summary="Return sheet/row data for an Excel file as JSON (for in-browser preview)",
)
async def spreadsheet_preview(
    file_id: int,
    max_rows_per_sheet: int = Query(default=500, ge=1, le=2000),
    db: AsyncSession = Depends(get_db),
):
    """
    Parse an XLSX/XLS file from S3 and return its content as JSON.

    Response shape::

        {
          "file_id": 123,
          "filename": "schedules.xlsx",
          "sheets": [
            {
              "name": "P&L",
              "rows": [
                {"row": 1, "cells": ["Item", "FY24", "FY23"]},
                ...
              ],
              "truncated": false
            }
          ]
        }

    Returns 409 if the file is not an Excel workbook.
    """
    from src.services.document_text import is_xlsx_filename

    obj = await PortfolioService.get_file(db, file_id)
    if not obj:
        raise HTTPException(status_code=404, detail="File not found")
    if not obj.storage_uri:
        raise HTTPException(status_code=409, detail="File has no storage URI")
    if not is_xlsx_filename(obj.filename or ""):
        raise HTTPException(
            status_code=409,
            detail="This file is not an Excel workbook. Spreadsheet preview is only available for .xlsx/.xls files.",
        )
    try:
        data = await asyncio.to_thread(download_storage_uri, obj.storage_uri)
    except Exception as exc:
        raise HTTPException(status_code=502, detail="Could not fetch file from storage") from exc

    try:
        import openpyxl  # type: ignore[import-untyped]
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl is not installed on this server")

    def _parse() -> list[dict]:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        sheets = []
        for sheet_name in wb.sheetnames:
            try:
                ws = wb[sheet_name]
            except Exception:
                continue
            rows_out = []
            truncated = False
            for row_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
                if not any(c is not None and str(c).strip() for c in row):
                    continue
                rows_out.append({
                    "row": row_idx,
                    "cells": [None if c is None else str(c) for c in row],
                })
                if len(rows_out) >= max_rows_per_sheet:
                    truncated = True
                    break
            sheets.append({"name": sheet_name, "rows": rows_out, "truncated": truncated})
        wb.close()
        return sheets

    try:
        sheets = await asyncio.to_thread(_parse)
    except Exception as exc:
        logger.exception("spreadsheet preview parse failed file_id=%s", file_id)
        raise HTTPException(status_code=422, detail="Could not parse Excel file") from exc

    return {"file_id": file_id, "filename": obj.filename, "sheets": sheets}


@router.get("/files/{file_id}", response_model=FileRead)
async def get_file(
    file_id: int,
    db: AsyncSession = Depends(get_db),
):
    obj = await PortfolioService.get_file(db, file_id)
    if not obj:
        raise HTTPException(status_code=404, detail="File not found")
    return _file_read_from_orm(obj)


@router.patch("/files/{file_id}")
async def patch_file(
    file_id: int,
    payload: FilePatch,
    db: AsyncSession = Depends(get_db),
):
    try:
        obj = await PortfolioService.patch_file(db, file_id, payload)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e)) from None
    if not obj:
        raise HTTPException(status_code=404, detail="File not found")
    result = _file_read_from_orm(obj).model_dump()
    # Warn when re-tagging a file to an entity+cycle that already has other files.
    entity_id_in_patch = payload.model_fields_set and "entity_id" in payload.model_fields_set
    if entity_id_in_patch and obj.entity_id is not None:
        warning = await AuditService(db)._check_duplicate_files(
            obj.entity_id, obj.review_cycle_id, exclude_file_id=obj.id
        )
        if warning:
            result.update(warning)
    return result


@router.delete("/files/{file_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_file(
    file_id: int,
    db: AsyncSession = Depends(get_db),
):
    ok = await PortfolioService.delete_file(db, file_id)
    if not ok:
        raise HTTPException(status_code=404, detail="File not found")
    return None


# ---- Reconciliation export ----

@router.get("/financial-metric-reconciliation/export/in-review", tags=["Reconciliation"])
async def export_in_review_reconciliation(
    review_cycle_id: Optional[str] = Query(default=None, description="Restrict to one review cycle"),
    db: AsyncSession = Depends(get_db),
):
    """Stream XLSX of all financial metric reconciliation rows for In Review companies."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl is not installed")

    from sqlalchemy import exists, func, select
    from src.db.models import Entity, FinancialMetricReconciliation, ManualReconciliationQuery, PortfolioCompany, ReviewCycle
    from src.schema.portfolio import ENTITY_ACTIVE_REVIEW_STATUSES

    rc_rows = (await db.execute(select(ReviewCycle))).scalars().all()
    rc_labels: dict[str, str] = {rc.id: (rc.name or rc.id) for rc in rc_rows}

    # "In Review" companies = companies with at least one entity actively under review
    # (state lives on entities.status now).
    pc_stmt = select(PortfolioCompany).where(
        exists(
            select(1).where(
                Entity.portfolio_company_id == PortfolioCompany.id,
                Entity.status.in_(ENTITY_ACTIVE_REVIEW_STATUSES),
            )
        )
    )
    if review_cycle_id:
        pc_stmt = pc_stmt.where(PortfolioCompany.review_cycle_id == review_cycle_id)
    companies = (await db.execute(pc_stmt)).scalars().all()
    company_map: dict[int, PortfolioCompany] = {pc.id: pc for pc in companies}

    def _empty_xlsx(sheet_title: str):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = sheet_title
        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return buf

    if not company_map:
        buf = _empty_xlsx("Reconciliation")
        fname = f"in-review-reconciliation-{date.today()}.xlsx"
        return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                 headers={"Content-Disposition": f'attachment; filename="{fname}"'})

    company_ids = list(company_map.keys())
    ent_rows = (await db.execute(select(Entity).where(Entity.portfolio_company_id.in_(company_ids)))).scalars().all()
    entity_map: dict[int, str] = {e.id: e.name for e in ent_rows}

    recon_rows = (
        await db.execute(
            select(FinancialMetricReconciliation)
            .where(FinancialMetricReconciliation.portfolio_company_id.in_(company_ids))
            .order_by(
                FinancialMetricReconciliation.portfolio_company_id,
                FinancialMetricReconciliation.entity_id,
                FinancialMetricReconciliation.review_cycle,
                FinancialMetricReconciliation.metric_key,
            )
        )
    ).scalars().all()

    from src.services.financial_reconciliation import fetch_canonical_snowflake_map, resolve_mis_for_recon_row

    recon_rows = list(recon_rows)
    sf_map = await fetch_canonical_snowflake_map(
        db,
        {(r.portfolio_company_id, (r.review_cycle or "").strip()) for r in recon_rows if (r.review_cycle or "").strip()},
    )

    manual_rows = (
        await db.execute(
            select(ManualReconciliationQuery)
            .where(ManualReconciliationQuery.portfolio_company_id.in_(company_ids))
            .order_by(ManualReconciliationQuery.portfolio_company_id, ManualReconciliationQuery.id)
        )
    ).scalars().all()

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Financial Reconciliation"

    hdr_font = Font(bold=True, color="FFFFFF")
    hdr_fill = PatternFill("solid", fgColor="2563EB")
    hdr_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

    recon_headers = [
        "ID", "Company Name", "Company ID", "Review Cycle", "Fund", "Investment Lead",
        "Contact Name", "Contact Email", "Entity Name", "Entity ID", "Metric",
        "MIS Amount", "MIS Currency", "AFS Amount", "AFS Currency",
        "Category", "Status", "Enabled", "Flagged", "Discrepancy Text",
        "Company Response", "Reviewer Remarks", "Variance Category",
        "Created At", "Updated At",
    ]
    for ci, h in enumerate(recon_headers, 1):
        cell = ws.cell(row=1, column=ci, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = hdr_align
    ws.row_dimensions[1].height = 28

    for ri, r in enumerate(recon_rows, 2):
        pc = company_map.get(r.portfolio_company_id)
        cycle_label = rc_labels.get(pc.review_cycle_id or "", pc.review_cycle_id or "—") if pc else "—"
        ent_name = entity_map.get(r.entity_id, "—") if r.entity_id else "—"
        sf = sf_map.get((r.portfolio_company_id, (r.review_cycle or "").strip()))
        mis_amount, mis_currency = resolve_mis_for_recon_row(r, sf)
        ws.append([
            r.id,
            pc.name if pc else "—",
            pc.company_id if pc else "—",
            cycle_label,
            pc.fund or "—" if pc else "—",
            pc.investment_lead or "—" if pc else "—",
            pc.contact_name or "—" if pc else "—",
            pc.contact_email_id or "—" if pc else "—",
            ent_name,
            r.entity_id,
            r.metric_key,
            mis_amount,
            mis_currency or "",
            float(r.afs_amount) if r.afs_amount is not None else None,
            r.afs_currency or "",
            r.category,
            r.status or "—",
            "Yes" if r.enable else "No",
            "Yes" if r.flagged else "No",
            r.discrepency_text or "",
            r.company_response or "",
            r.reviewer_remarks or "",
            r.variance_category or "",
            r.created_at.isoformat() if r.created_at else "",
            r.updated_at.isoformat() if r.updated_at else "",
        ])

    ws.freeze_panes = "A2"

    # Second sheet: manual queries
    ws2 = wb.create_sheet("Manual Queries")
    mq_headers = [
        "ID", "Company Name", "Entity Name", "Query Text",
        "Status", "Enabled", "Flagged", "Company Response", "Reviewer Remarks",
        "Created At",
    ]
    for ci, h in enumerate(mq_headers, 1):
        cell = ws2.cell(row=1, column=ci, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = hdr_align
    for ri, mq in enumerate(manual_rows, 2):
        pc = company_map.get(mq.portfolio_company_id)
        ent_name = entity_map.get(mq.entity_id, "—") if mq.entity_id else "—"
        ws2.append([
            mq.id,
            pc.name if pc else "—",
            ent_name,
            mq.discrepency_text or "",
            mq.status or "—",
            "Yes" if mq.enable else "No",
            "Yes" if mq.flagged else "No",
            mq.company_response or "",
            mq.reviewer_remarks or "",
            mq.created_at.isoformat() if mq.created_at else "",
        ])
    ws2.freeze_panes = "A2"

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    safe_rc = (review_cycle_id or "").replace("/", "-").replace(" ", "_")
    fname = f"in-review-reconciliation-{safe_rc + '-' if safe_rc else ''}{date.today()}.xlsx"
    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{fname}"'},
    )
